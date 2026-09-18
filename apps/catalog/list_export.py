"""FEAT-096 — export Excel de la liste catalogue actuellement filtrée.

Ce n'est pas l'export d'inventaire de FEAT-078 (toutes les colonnes d'import,
tout le fonds). Ici on reprend **ce que la page Catalogue montre** — notices
ou exemplaires, avec les filtres courants — et on ajoute la colonne
emplacement, absente de l'écran.
"""
from __future__ import annotations

import io

from django.utils import timezone
from django.utils.translation import gettext as _

_RECORD_WIDTHS = [44, 30, 24, 8, 20]
_ITEM_WIDTHS = [44, 30, 24, 16, 18, 24, 12]


def _authors(record) -> str:
    return ", ".join(a.full_name for a in record.authors.all())


def _record_locations(record) -> str:
    codes = {
        item.location.code
        for item in record.items.all()
        if item.location_id and item.location
    }
    return ", ".join(sorted(codes))


def _sorted_records(params):
    from .views import filtered_records

    records, relevance = filtered_records(params)
    records = records.prefetch_related("items__location")
    rows = list(records)
    if relevance is not None:
        rows.sort(key=lambda record: relevance.get(record.pk, 1 << 30))
    else:
        rows.sort(key=lambda record: record.title.lower())
    return rows


def _sorted_items(params):
    from .views import filtered_items, filtered_records

    records, relevance = filtered_records(params)
    items = filtered_items(params, records)
    if relevance is not None:
        rows = list(items)
        rows.sort(
            key=lambda item: (relevance.get(item.record_id, 1 << 30), item.internal_id)
        )
        return rows
    return list(items.order_by("record__title", "internal_id"))


def is_items_mode(params) -> bool:
    return (params.get("mode") or "") == "items"


def list_export_filename(params) -> str:
    kind = "exemplaires" if is_items_mode(params) else "notices"
    return f"catalogue-{kind}-{timezone.localdate().isoformat()}.xlsx"


def record_headers() -> list[str]:
    return [
        _("Titre"),
        _("Auteur(s)"),
        _("Classification"),
        _("Ex."),
        _("Emplacement"),
    ]


def item_headers() -> list[str]:
    return [
        _("Titre"),
        _("Auteur(s)"),
        _("Classification"),
        _("Code Ofelia"),
        _("Code Ofelia externe"),
        _("Provenance"),
        _("Emplacement"),
    ]


def record_row(record) -> list:
    return [
        record.title,
        _authors(record),
        record.category.name if record.category else "",
        len(record.items.all()),
        _record_locations(record),
    ]


def item_row(item) -> list:
    record = item.record
    return [
        record.title,
        _authors(record),
        record.category.name if record.category else "",
        item.ean13,
        item.external_code or "",
        str(item.provenance) if item.provenance else "",
        item.location.code if item.location else "",
    ]


def _workbook(sheet_title: str, headers: list[str], widths: list[int], rows) -> bytes:
    import openpyxl
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet(title=sheet_title[:31] or "Catalogue")
    bold = Font(bold=True)
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"

    header_row = []
    for name in headers:
        cell = WriteOnlyCell(ws, value=name)
        cell.font = bold
        header_row.append(cell)
    ws.append(header_row)
    for row in rows:
        ws.append(list(row))

    buffer = io.BytesIO()
    wb.save(buffer)
    wb.close()
    return buffer.getvalue()


def build_list_workbook(params) -> bytes:
    """Classeur de la recherche courante, une ligne par notice ou exemplaire."""
    if is_items_mode(params):
        return _workbook(
            _("Exemplaires"),
            item_headers(),
            _ITEM_WIDTHS,
            (item_row(item) for item in _sorted_items(params)),
        )
    return _workbook(
        _("Notices"),
        record_headers(),
        _RECORD_WIDTHS,
        (record_row(record) for record in _sorted_records(params)),
    )
