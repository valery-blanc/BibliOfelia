"""FEAT-093 — les corrections demandées par Val après le premier essai.

Camemberts étiquetés sur les tranches, boutons de période, menu des
sous-rapports, exports par sous-rapport, graphes appariés aux tableaux, et les
données brutes rapatriées dans les écrans qui les montrent.
"""
from __future__ import annotations

from datetime import date, timedelta
from decimal import Decimal

import pytest
from django.urls import reverse
from django.utils import timezone

from apps.accounts.models import Role
from apps.catalog.models import AcquisitionSource, BibliographicRecord, Category, Item
from apps.closing.models import (
    ActivityEntry,
    ActivityType,
    AnimationAttendance,
    AnimationSession,
    AnimationType,
)
from apps.finance.models import FeeKind, Invoice, InvoiceLine, Payment
from apps.loans.models import Loan, LoanStatus
from apps.members.models import CardRenewal, Member, MemberCategory, MemberFamilyMember
from apps.reports import builders, periods
from apps.reports.pages import Chart, Series

pytestmark = pytest.mark.django_db

ALL_SLUGS = [spec.slug for spec in builders.REPORTS]


@pytest.fixture
def librarian(django_user_model):
    return django_user_model.objects.create_user(
        username="bib2", password="x", role=Role.LIBRARIAN,
        first_name="Katty", last_name="Roussel",
    )


@pytest.fixture
def client_librarian(client, librarian):
    client.force_login(librarian)
    return client


@pytest.fixture
def dataset(librarian):
    """Le même petit jeu que la suite principale : il doit toucher tous les écrans."""
    today = date.today()
    category = Category.objects.create(code="ALB", name="Albums")
    mcat = MemberCategory.objects.create(
        code="AD", name="Adulte", card_validity_months=12,
        membership_fee=Decimal("20.00"),
    )
    member = Member.objects.create(
        first_name="Ada", last_name="Lovelace", category=mcat,
        birth_date=date(1990, 5, 3), address_postal_code="1290",
        address_city="Versoix", contact_phone="0221234567",
        spoken_languages=["fr", "en"],
        expiration_date=today + timedelta(days=15),
    )
    MemberFamilyMember.objects.create(
        member=member, first_name="Nils", birth_year=today.year - 7
    )
    CardRenewal.objects.create(
        member=member, renewed_on=today - timedelta(days=10),
        previous_expiration=today, new_expiration=today + timedelta(days=365),
    )
    Member.objects.create(
        first_name="Perdu", last_name="Parti", category=mcat,
        expiration_date=today - timedelta(days=30),
    )

    record = BibliographicRecord.objects.create(
        title="Fondation", language="fr", category=category, document_type="book"
    )
    item = Item.objects.create(
        record=record, acquisition_source=AcquisitionSource.DONATION, donor="Paroisse"
    )
    Loan.objects.create(
        item=item, member=member, due_date=today - timedelta(days=9),
        status=LoanStatus.OVERDUE,
    )
    Loan.objects.create(
        item=item, member=member, due_date=today - timedelta(days=30),
        return_date=timezone.now(), status=LoanStatus.RETURNED,
    )

    animation = AnimationType.objects.create(label="Bibliothèque")
    session = AnimationSession.objects.create(
        animation_type=animation, presenter=librarian, minutes=120,
        non_member_adults=2, non_member_children=5,
    )
    AnimationAttendance.objects.create(session=session, member=member)
    ActivityEntry.objects.create(
        user=librarian, activity_type=ActivityType.objects.create(label="Ouverture"),
        minutes=180,
    )

    invoice = Invoice.objects.create(member=member, due_date=today + timedelta(days=30))
    InvoiceLine.objects.create(
        invoice=invoice, kind=FeeKind.MEMBERSHIP, label="Cotisation",
        amount=Decimal("20.00"),
    )
    invoice.recompute()
    Payment.objects.create(invoice=invoice, amount=Decimal("20.00"))
    return {"member": member, "item": item, "category": category}


def _period():
    return periods.Period(
        date.today() - timedelta(days=60), date.today(), "test", "custom"
    )


# ── Camemberts : les libellés sur les tranches ─────────────────────────────


def test_every_slice_gets_its_label_outside_with_a_leader_line():
    """« Adultes Documentaire 35 (34 %) » se lit à côté de sa part, reliée par
    un trait — comme dans le PDF.

    Écrire dans la tranche ne marchait pas : un texte blanc posé sur une part
    claire, ou débordant sur le fond blanc de la page, devenait invisible
    (retour de Val). À l'extérieur, le texte est toujours en encre foncée sur
    fond blanc, quelle que soit la couleur de la part.
    """
    from apps.reports import charts

    chart = Chart(
        title="t", kind="pie",
        labels=["Adultes Documentaire", "Jeunesse"], values=[35, 68],
    )
    svg = charts.pie_svg(chart)
    assert "Adultes Documentaire" in svg
    assert "34 %" in svg  # 35 / 103
    assert "66 %" in svg
    # Une ligne de rappel par part, et aucun texte blanc posé sur le dessin.
    assert svg.count("<polyline") == 2
    assert 'fill="#FFFFFF"' not in svg


def test_a_tiny_slice_gets_an_outside_label_with_a_leader_line():
    """Une part minuscule doit rester lisible : son texte est déporté."""
    from apps.reports import charts

    svg = charts.pie_svg(
        Chart(title="t", kind="pie", labels=["Gros", "Miette"], values=[500, 3])
    )
    assert "<polyline" in svg
    assert "Miette" in svg


def test_crowded_labels_do_not_overlap():
    """Huit parts, donc huit étiquettes : elles s'empilent, et le dessin
    s'agrandit assez pour toutes les contenir."""
    from apps.reports import charts

    chart = Chart(
        title="t", kind="pie",
        labels=[f"Rayon {i}" for i in range(8)],
        values=[100, 50, 25, 12, 6, 3, 2, 1],
    )
    svg = charts.pie_svg(chart)
    assert svg.count("<polyline") == 8
    # Deux étiquettes voisines sont séparées d'au moins une hauteur de ligne.
    import re

    ys = sorted(float(m) for m in re.findall(r'<text[^>]*y="([\d.]+)"', svg))
    side_gaps = [b - a for a, b in zip(ys, ys[1:]) if b - a > 0.1]
    assert min(side_gaps) >= charts.LABEL_STEP - 0.5


def test_slice_label_wording():
    from apps.reports.charts import slice_label

    assert slice_label("Adultes Documentaire", 35, 103) == "Adultes Documentaire 35 (34 %)"


def test_a_single_slice_pie_is_still_drawn():
    """Un arc dont le début et la fin coïncident ne trace rien : une catégorie
    unique à 100 % donnerait un camembert vide."""
    from apps.reports import charts

    svg = charts.pie_svg(Chart(title="t", kind="pie", labels=["Tout"], values=[7]))
    assert svg.count("<path") == 2  # l'anneau complet, en deux moitiés
    assert "100 %" in svg


# ── Courbes ────────────────────────────────────────────────────────────────


def test_line_chart_draws_one_polyline_per_series():
    from apps.reports import charts

    chart = Chart(
        title="t", kind="line", labels=["jan", "fév", "mar"],
        series=[
            Series("Nouveaux", [3, 5, 2]),
            Series("Réinscriptions", [1, 0, 4]),
            Series("Usagers perdus", [0, 2, 1]),
        ],
    )
    svg = charts.line_svg(chart)
    # Trois courbes ; les repères horizontaux sont des <line>, pas des polylignes.
    assert svg.count("<polyline") == 3
    for label in ("Nouveaux", "Réinscriptions", "Usagers perdus"):
        assert label in svg


def test_an_all_zero_line_chart_counts_as_empty():
    chart = Chart(
        title="t", kind="line", labels=["jan"],
        series=[Series("Nouveaux", [0]), Series("Perdues", [0])],
    )
    assert chart.is_empty


# ── Tableaux appariés à leur graphe ────────────────────────────────────────


@pytest.mark.parametrize(
    "slug,title",
    [
        ("members", "Qui arrive, qui reste, qui part"),
        ("members", "D'où viennent les usagers"),
        ("collection", "Les rayons : ce qu'on a et ce qui sort"),
        ("team", "Le temps par nature de travail"),
        ("attendance", "Ce qui a été organisé"),
        ("attendance", "Les présences par âge et par animation"),
    ],
)
def test_these_tables_carry_the_chart_of_their_own_figures(dataset, slug, title):
    page = builders.get(slug).builder(_period(), {})
    table = next(t for t in page.tables if t.title == title)
    assert table.chart is not None


def test_the_movement_chart_has_the_three_expected_lines(dataset):
    page = builders.get("members").builder(_period(), {})
    table = next(t for t in page.tables if t.title == "Qui arrive, qui reste, qui part")
    assert table.chart.kind == "line"
    assert [s.label for s in table.chart.series] == [
        "Nouveaux", "Réinscriptions", "Usagers perdus"
    ]


def test_the_locality_chart_is_a_pie(dataset):
    page = builders.get("members").builder(_period(), {})
    table = next(t for t in page.tables if t.title == "D'où viennent les usagers")
    assert table.chart.kind == "pie"


def test_the_team_chart_is_a_pie(dataset):
    """« À quoi passe-t-on le temps ? » est un partage d'un tout : camembert,
    pas barres (retour Val)."""
    page = builders.get("team").builder(_period(), {})
    table = next(t for t in page.tables if t.title == "Le temps par nature de travail")
    assert table.chart.kind == "pie"


def test_team_no_longer_says_the_same_thing_twice(dataset):
    """Les deux sous-rapports « par nature de travail » ont fusionné."""
    page = builders.get("team").builder(_period(), {})
    titles = [b.title for b in page.blocks]
    assert titles.count("Le temps par nature de travail") == 1
    assert "Les heures par nature de travail" not in titles


def test_animation_hours_are_not_multiplied_by_the_number_of_attendees(dataset):
    """La séance dure 120 minutes et compte un présent : deux heures, pas plus."""
    page = builders.get("attendance").builder(_period(), {})
    table = next(t for t in page.tables if t.title == "Ce qui a été organisé")
    assert table.rows[0][-1] == "2,0"


def test_the_attendance_pie_counts_people_not_sessions(dataset):
    """Une séance, un membre présent, sept non-membres : le camembert partage
    huit présences, pas une séance."""
    page = builders.get("attendance").builder(_period(), {})
    table = next(t for t in page.tables if t.title == "Ce qui a été organisé")
    assert table.chart.values == [8]


# Le bilan de la vue d'ensemble a changé deux fois : il est décrit par
# `test_report_final_touches.py`, qui porte sa forme définitive — deux colonnes,
# pas de graphe.


# ── Données brutes rapatriées ──────────────────────────────────────────────


def test_raw_data_moved_into_the_screens_that_show_them(dataset):
    period = periods.default_period()
    collection_titles = [t.title for t in builders.get("collection").builder(period, {}).tables]
    loans_titles = [t.title for t in builders.get("loans").builder(period, {}).tables]
    assert "Le catalogue complet" in collection_titles
    assert "Prêts et réservations en cours" in loans_titles
    assert "Le détail des prêts de la période" in loans_titles


def test_the_hub_no_longer_lists_raw_exports(client_librarian, dataset):
    html = client_librarian.get(reverse("reports:index")).content.decode()
    assert "Sortir les données brutes" not in html
    # Les CSV restent atteignables, mais depuis l'écran qui montre les données.
    page = client_librarian.get(reverse("reports:view", kwargs={"slug": "collection"}))
    assert reverse("reports:catalog_csv") in page.content.decode()


def test_the_csv_exports_still_work(client_librarian, dataset):
    assert client_librarian.get(reverse("reports:catalog_csv")).status_code == 200
    assert client_librarian.get(
        reverse("reports:active_loans_reservations_csv")
    ).status_code == 200


# ── Exports par sous-rapport ───────────────────────────────────────────────


@pytest.mark.parametrize("slug", ALL_SLUGS)
def test_each_sub_report_exports_on_its_own(client_librarian, dataset, slug):
    """On doit pouvoir imprimer la seule liste des impayés sans sortir les huit
    tableaux de l'écran."""
    view = reverse("reports:view", kwargs={"slug": slug})
    page = client_librarian.get(view).context["page"]
    assert page.blocks, f"{slug} n'a aucun sous-rapport"
    key = page.blocks[0].key

    pdf = client_librarian.get(
        reverse("reports:pdf", kwargs={"slug": slug}) + f"?block={key}"
    )
    assert pdf.status_code == 200
    assert pdf.content[:4] == b"%PDF"
    assert key in pdf["Content-Disposition"]

    xlsx = client_librarian.get(
        reverse("reports:xlsx", kwargs={"slug": slug}) + f"?block={key}"
    )
    assert xlsx.status_code == 200
    assert xlsx.content[:2] == b"PK"


def test_an_unknown_block_key_falls_back_to_the_whole_screen(client_librarian, dataset):
    """Un lien périmé rend l'écran entier, jamais une erreur."""
    resp = client_librarian.get(
        reverse("reports:pdf", kwargs={"slug": "collection"}) + "?block=nexistepas"
    )
    assert resp.status_code == 200
    assert resp.content[:4] == b"%PDF"


def test_block_keys_are_stable_across_calls(dataset):
    """La clé vient du titre, pas de la position : un lien d'export partagé ne
    doit pas pointer sur un autre tableau après réordonnancement."""
    period = periods.default_period()
    first = {t.title: t.key for t in builders.get("collection").builder(period, {}).tables}
    second = {t.title: t.key for t in builders.get("collection").builder(period, {}).tables}
    assert first == second


def test_a_sub_report_export_keeps_the_period(client_librarian, dataset):
    view = reverse("reports:view", kwargs={"slug": "loans"}) + "?p=year"
    page = client_librarian.get(view).context["page"]
    key = page.tables[0].key
    resp = client_librarian.get(
        reverse("reports:pdf", kwargs={"slug": "loans"}) + f"?block={key}&p=year"
    )
    assert resp.status_code == 200
    assert str(date.today().year) in resp["Content-Disposition"]


# ── Boutons et menu ────────────────────────────────────────────────────────


def test_period_and_filter_choices_are_rendered_as_buttons(client_librarian, dataset):
    """Val veut des boutons, avec une couleur distincte pour celui qui est actif."""
    html = client_librarian.get(
        reverse("reports:view", kwargs={"slug": "loans"}) + "?p=year"
    ).content.decode()
    assert "chip-btn" in html
    assert "chip-btn--on" in html

    overdue = client_librarian.get(
        reverse("reports:view", kwargs={"slug": "overdue"}) + "?threshold=14"
    ).content.decode()
    assert "chip-btn--on" in overdue


def test_every_screen_offers_its_sub_report_menu(client_librarian, dataset):
    html = client_librarian.get(
        reverse("reports:view", kwargs={"slug": "collection"})
    ).content.decode()
    assert "report-toc" in html
    page = client_librarian.get(
        reverse("reports:view", kwargs={"slug": "collection"})
    ).context["page"]
    for block in page.blocks:
        assert f'href="#{block.key}"' in html


def test_the_hub_carries_a_colour_for_each_report(client_librarian, dataset):
    html = client_librarian.get(reverse("reports:index")).content.decode()
    assert "--lr-bg:var(--blush-light)" in html  # les retards
    assert "--lr-bg:var(--forest-light)" in html  # le fonds


# ── Excel ──────────────────────────────────────────────────────────────────


def test_excel_puts_a_paired_chart_on_the_table_sheet(dataset):
    """Tableau et graphe des mêmes données restent ensemble dans le tableur,
    comme à l'écran."""
    import io

    import openpyxl

    from apps.reports.excel import render_report_xlsx

    page = builders.get("members").builder(_period(), {})
    table = next(t for t in page.tables if t.title == "D'où viennent les usagers")
    workbook = openpyxl.load_workbook(io.BytesIO(render_report_xlsx(page, table)))
    sheet = workbook[workbook.sheetnames[-1]]
    assert sheet._charts, "le graphe apparié doit être sur la feuille du tableau"


def test_excel_pie_labels_show_name_value_and_percent(dataset):
    """Le camembert du fichier Excel porte les mêmes étiquettes qu'à l'écran."""
    import io

    import openpyxl

    from apps.reports.excel import render_report_xlsx

    page = builders.get("members").builder(_period(), {})
    table = next(t for t in page.tables if t.title == "D'où viennent les usagers")
    workbook = openpyxl.load_workbook(io.BytesIO(render_report_xlsx(page, table)))
    chart = workbook[workbook.sheetnames[-1]]._charts[0]
    assert chart.dataLabels.showCatName
    assert chart.dataLabels.showVal
    assert chart.dataLabels.showPercent


def test_excel_of_a_line_chart_holds_one_series_per_line(dataset):
    import io

    import openpyxl

    from apps.reports.excel import render_report_xlsx

    page = builders.get("members").builder(_period(), {})
    table = next(t for t in page.tables if t.title == "Qui arrive, qui reste, qui part")
    workbook = openpyxl.load_workbook(io.BytesIO(render_report_xlsx(page, table)))
    chart = workbook[workbook.sheetnames[-1]]._charts[0]
    assert len(chart.series) == 3
