"""FEAT-093 — les dix écrans de rapport, et leurs deux exports.

Le test central est `test_every_report_has_its_two_exports` : il parcourt le
registre, et non une liste écrite à la main. Un écran ajouté sans son PDF ou son
fichier Excel fait donc échouer la suite, ce qui est exactement la règle posée
par Val — pas d'écran sans ses deux exports.
"""
from __future__ import annotations

from datetime import date, timedelta
from decimal import Decimal

import pytest
from django.urls import reverse
from django.utils import timezone

from apps.accounts.models import Role
from apps.catalog.models import (
    AcquisitionSource,
    BibliographicRecord,
    Category,
    Item,
    ItemStatus,
)
from apps.closing.models import (
    ActivityEntry,
    ActivityType,
    AnimationAttendance,
    AnimationSession,
    AnimationType,
)
from apps.finance.models import (
    CashDirection,
    CashMovement,
    FeeKind,
    Invoice,
    InvoiceLine,
    Payment,
)
from apps.loans.models import Loan, LoanStatus, Reservation, ReservationStatus
from apps.members.models import CardRenewal, Member, MemberCategory, MemberFamilyMember
from apps.reports import builders, periods
from apps.reports.pages import Chart, Table

pytestmark = pytest.mark.django_db


@pytest.fixture
def librarian(django_user_model):
    return django_user_model.objects.create_user(
        username="bib", password="x", role=Role.LIBRARIAN,
        first_name="Katty", last_name="Roussel",
    )


@pytest.fixture
def client_librarian(client, librarian):
    client.force_login(librarian)
    return client


@pytest.fixture
def dataset(librarian):
    """Un petit jeu de données qui touche **tous** les écrans.

    Les rapports agrègent sept applications : une base vide ne prouverait que
    l'absence d'erreur de syntaxe.
    """
    today = date.today()
    category = Category.objects.create(code="ALB", name="Albums")
    mcat = MemberCategory.objects.create(
        code="AD", name="Adulte", card_validity_months=12,
        membership_fee=Decimal("20.00"),
    )
    member = Member.objects.create(
        first_name="Ada", last_name="Lovelace", category=mcat,
        birth_date=date(1990, 5, 3), address_postal_code="1290",
        address_city="Versoix", contact_phone="0221234567",
        spoken_languages=["fr", "en"],
        expiration_date=today + timedelta(days=15),
    )
    MemberFamilyMember.objects.create(
        member=member, first_name="Nils", birth_year=today.year - 7
    )
    CardRenewal.objects.create(
        member=member, renewed_on=today - timedelta(days=10),
        previous_expiration=today, new_expiration=today + timedelta(days=365),
    )
    # Un usager dont la carte a expiré et n'a pas été renouvelée : « perdu ».
    Member.objects.create(
        first_name="Perdu", last_name="Parti", category=mcat,
        expiration_date=today - timedelta(days=30),
    )

    record = BibliographicRecord.objects.create(
        title="Fondation", language="fr", category=category, document_type="book"
    )
    item = Item.objects.create(
        record=record, acquisition_source=AcquisitionSource.DONATION, donor="Paroisse"
    )
    Item.objects.create(record=record, status=ItemStatus.LOST)

    loan = Loan.objects.create(
        item=item, member=member, due_date=today - timedelta(days=9),
        status=LoanStatus.OVERDUE,
    )
    Loan.objects.create(
        item=item, member=member, due_date=today - timedelta(days=30),
        return_date=timezone.now(), status=LoanStatus.RETURNED,
    )
    Reservation.objects.create(
        record=record, member=member, expires_at=today + timedelta(days=7),
        status=ReservationStatus.READY_FOR_PICKUP, ready_since=today,
        fulfilled_by_item=item,
    )

    animation = AnimationType.objects.create(label="Bibliothèque")
    session = AnimationSession.objects.create(
        animation_type=animation, presenter=librarian, minutes=120,
        non_member_adults=2, non_member_children=5,
    )
    AnimationAttendance.objects.create(session=session, member=member)
    ActivityEntry.objects.create(
        user=librarian, activity_type=ActivityType.objects.create(label="Ouverture"),
        minutes=180,
    )

    invoice = Invoice.objects.create(member=member, due_date=today + timedelta(days=30))
    InvoiceLine.objects.create(
        invoice=invoice, kind=FeeKind.MEMBERSHIP, label="Cotisation",
        amount=Decimal("20.00"),
    )
    invoice.recompute()
    Payment.objects.create(invoice=invoice, amount=Decimal("20.00"))
    CashMovement.objects.create(
        direction=CashDirection.OUT, amount=Decimal("5.00"), label="Timbres"
    )
    return {"member": member, "item": item, "loan": loan, "category": category}


ALL_SLUGS = [spec.slug for spec in builders.REPORTS]


@pytest.mark.parametrize("slug", ALL_SLUGS)
def test_every_report_has_its_two_exports(client_librarian, dataset, slug):
    """Chaque écran s'affiche, s'imprime en PDF et s'enregistre en Excel."""
    html = client_librarian.get(reverse("reports:view", kwargs={"slug": slug}))
    assert html.status_code == 200

    pdf = client_librarian.get(reverse("reports:pdf", kwargs={"slug": slug}))
    assert pdf.status_code == 200
    assert pdf["Content-Type"] == "application/pdf"
    assert pdf.content[:4] == b"%PDF"

    xlsx = client_librarian.get(reverse("reports:xlsx", kwargs={"slug": slug}))
    assert xlsx.status_code == 200
    # Un .xlsx est une archive zip : le fichier commence par « PK ».
    assert xlsx.content[:2] == b"PK"


@pytest.mark.parametrize("slug", ALL_SLUGS)
def test_reports_survive_an_empty_database(client_librarian, slug):
    """Une bibliothèque qui vient d'être installée n'a rien à montrer, et cela
    ne doit produire ni erreur ni page blanche : des tableaux vides et leur
    phrase d'explication."""
    resp = client_librarian.get(reverse("reports:view", kwargs={"slug": slug}))
    assert resp.status_code == 200
    assert client_librarian.get(
        reverse("reports:pdf", kwargs={"slug": slug})
    ).status_code == 200


def test_unknown_report_is_a_404(client_librarian):
    assert client_librarian.get("/fr/reports/licorne/").status_code == 404


def test_hub_lists_the_work_counts(client_librarian, dataset):
    resp = client_librarian.get(reverse("reports:index"))
    assert resp.status_code == 200
    counts = {entry["spec"].slug: entry["count"] for entry in resp.context["work"]}
    assert counts["overdue"] == 1
    assert counts["pickup"] == 1


def test_readonly_may_read_a_report_and_its_exports(client, django_user_model, dataset):
    """Un membre du comité en lecture seule doit pouvoir sortir le PDF : c'est
    précisément la personne à qui la vue d'ensemble est destinée."""
    user = django_user_model.objects.create_user(
        username="comite", password="x", role=Role.READONLY
    )
    client.force_login(user)
    assert client.get(reverse("reports:view", kwargs={"slug": "overview"})).status_code == 200
    assert client.get(reverse("reports:pdf", kwargs={"slug": "overview"})).status_code == 200


# ── Le sélecteur de période ────────────────────────────────────────────────


def test_default_period_is_the_last_twelve_months():
    """Douze mois bornes comprises, soit 365 jours — pas 366.

    `aujourd'hui − 365 jours` donnerait une fenêtre d'un jour de trop et
    décalerait d'autant la période de comparaison.
    """
    period = periods.default_period(date(2026, 9, 9))
    assert period.start == date(2025, 9, 10)
    assert period.end == date(2026, 9, 9)
    assert period.days == 365


def test_the_four_shortcuts_carry_real_labels():
    labels = [p.label for p in periods.shortcuts(date(2026, 9, 9))]
    assert labels[1] == "2026"  # année en cours
    assert labels[3] == "2025"  # année précédente
    # Mois en cours et mois précédent portent un nom de mois, pas « précédent ».
    assert "2026" in labels[0] and "2026" in labels[2]


def test_a_named_period_covers_the_whole_month():
    period = periods.named_period("prev_month", date(2026, 3, 15))
    assert period.start == date(2026, 2, 1)
    assert period.end == date(2026, 2, 28)


def test_reversed_dates_fall_back_instead_of_erroring():
    period, warning = periods.parse(
        {"p": "custom", "start": "2026-09-09", "end": "2026-01-01"}, date(2026, 9, 9)
    )
    assert warning
    assert period.start == date(2025, 9, 10)


def test_an_unreadable_date_falls_back_instead_of_erroring():
    period, warning = periods.parse({"p": "custom", "start": "pas une date"}, date(2026, 9, 9))
    assert warning
    assert period.key == "custom"
# La période de comparaison a changé de définition (le même mois un an plus
# tôt, et non le mois précédent) : elle est décrite par
# `test_report_final_touches.py`.


def test_a_very_long_period_is_split_by_year_not_by_month():
    """Cinquante barres de mois ne se lisent pas : au-delà de trois ans on
    agrège par année."""
    long_period = periods.Period(date(2018, 1, 1), date(2026, 1, 1), "long", "custom")
    buckets = periods.months_between(long_period)
    assert [label for _s, _e, label in buckets] == [str(y) for y in range(2018, 2027)]


# ── Les définitions que la spec fixe ───────────────────────────────────────


def test_lost_members_are_expired_cards_that_were_not_renewed(dataset):
    period = periods.Period(
        date.today() - timedelta(days=60), date.today(), "test", "custom"
    )
    page = builders.get("members").builder(period, {})
    lost = next(k for k in page.kpis if "perdus" in str(k.label).lower())
    assert lost.value == "1"


def test_renewals_are_counted_from_the_card_renewal_trail(dataset):
    period = periods.Period(
        date.today() - timedelta(days=60), date.today(), "test", "custom"
    )
    page = builders.get("members").builder(period, {})
    renewals = next(k for k in page.kpis if "inscription" in str(k.label).lower())
    assert renewals.value == "1"


def test_renewing_a_card_writes_a_renewal(dataset, librarian):
    from apps.members.services import renew_card

    member = dataset["member"]
    before = CardRenewal.objects.count()
    renew_card(member, user=librarian, invoice=False)
    assert CardRenewal.objects.count() == before + 1
    # Un second clic le même jour ne change pas la date : pas de doublon.
    renew_card(member, user=librarian, invoice=False)
    assert CardRenewal.objects.count() == before + 1


def test_least_borrowed_only_lists_books_that_were_borrowed(dataset):
    """« Les moins empruntés » exclut les livres jamais sortis, qui sont comptés
    dans un tableau séparé — sinon la liste ne serait qu'une suite de zéros."""
    period = periods.Period(
        date.today() - timedelta(days=30), date.today(), "test", "custom"
    )
    page = builders.get("loans").builder(period, {})
    least = next(t for t in page.tables if "moins" in t.title.lower())
    for row in least.rows:
        assert row[-1] != "0"


def test_animation_hours_are_not_multiplied_by_the_number_of_attendees(dataset):
    """La séance dure 120 minutes et compte un présent : deux heures, pas plus.

    Compter les présences dans le même `annotate` que la somme des minutes
    multiplierait la durée par le nombre de présents.
    """
    period = periods.Period(
        date.today() - timedelta(days=30), date.today(), "test", "custom"
    )
    page = builders.get("attendance").builder(period, {})
    table = next(t for t in page.tables if t.title == "Ce qui a été organisé")
    assert table.rows[0][-1] == "2,0"


def test_attendance_ages_use_the_age_on_the_day_of_the_session():
    from apps.reports import format as fmt

    born = date(2016, 12, 31)
    assert fmt.age_on(born, date(2026, 6, 1)) == 9
    assert fmt.age_on(born, date(2027, 1, 1)) == 10
    assert fmt.age_bracket(fmt.age_on(born, date(2026, 6, 1))) == "6-10"
    assert fmt.age_bracket(None) == "âge inconnu"


def test_a_single_slice_pie_is_actually_drawn():
    """Un arc dont le début et la fin coïncident ne trace rien : une catégorie
    unique à 100 % donnerait un camembert vide."""
    from apps.reports import charts

    svg = charts.pie_svg(Chart(title="t", kind="pie", labels=["Tout"], values=[7]))
    assert svg.count("<path") == 2  # l'anneau complet est tracé en deux moitiés


def test_a_long_pie_tail_is_grouped_under_others():
    from apps.reports import charts

    chart = Chart(
        title="t", kind="pie",
        labels=[f"c{i}" for i in range(12)],
        values=[100 - i for i in range(12)],
    )
    grouped = charts.grouped_pairs(chart)
    assert len(grouped) == charts.MAX_SLICES
    assert grouped[-1][0] == "Autres"
    assert sum(v for _l, v in grouped) == sum(chart.values)


def test_old_addresses_still_lead_somewhere(client_librarian, dataset):
    """Le guide utilisateur et les favoris pointent encore sur les anciennes
    adresses. `overdue` et `inactive` ont gardé leur slug et s'affichent
    directement ; `reservations-pickup` a été renommée et redirige."""
    for old in ("/fr/reports/overdue/", "/fr/reports/inactive/"):
        assert client_librarian.get(old).status_code == 200

    resp = client_librarian.get("/fr/reports/reservations-pickup/")
    assert resp.status_code == 302
    assert resp["Location"] == reverse("reports:view", kwargs={"slug": "pickup"})


def test_the_csv_exports_are_still_there(client_librarian, dataset):
    """Les exports de données brutes (FEAT-040) ne sont pas des rapports et
    survivent à la refonte."""
    assert client_librarian.get(reverse("reports:catalog_csv")).status_code == 200


def test_export_links_carry_the_period(client_librarian, dataset):
    resp = client_librarian.get(
        reverse("reports:view", kwargs={"slug": "loans"}) + "?p=year"
    )
    assert resp.context["page"].query == "p=year"


def test_a_filter_survives_a_change_of_period(client_librarian, dataset):
    category = dataset["category"].code
    resp = client_librarian.get(
        reverse("reports:view", kwargs={"slug": "loans"})
        + f"?p=year&category={category}"
    )
    assert resp.context["filter_query"] == f"&category={category}"


def test_excel_holds_one_sheet_per_block(dataset):
    import io

    import openpyxl

    from apps.reports.excel import render_report_xlsx

    period = periods.default_period()
    page = builders.get("collection").builder(period, {})
    workbook = openpyxl.load_workbook(io.BytesIO(render_report_xlsx(page)))
    # Un onglet « Résumé » plus un onglet par bloc affiché à l'écran.
    assert len(workbook.sheetnames) == len(page.blocks) + 1
