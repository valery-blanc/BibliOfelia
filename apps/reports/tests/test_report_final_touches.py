"""FEAT-093 — les deux dernières demandes de Val (`temp.txt`, 2026-09-09).

1. « Les rayons : ce qu'on a et ce qui sort » troque son camembert pour un
   **double histogramme** : livres et prêts côte à côte, rayon par rayon.
2. « Le bilan en un tableau » retrouve sa colonne de **période précédente**,
   redéfinie — la même période un an plus tôt, sauf pour les dates libres — et
   perd son graphe.
"""
from __future__ import annotations

import io
from datetime import date, timedelta
from decimal import Decimal

import pytest

from apps.accounts.models import Role
from apps.catalog.models import BibliographicRecord, Category, Item
from apps.loans.models import Loan, LoanStatus
from apps.members.models import Member, MemberCategory
from apps.reports import builders, charts, periods
from apps.reports.excel import render_report_xlsx
from apps.reports.pages import Chart, Series
from apps.reports.pdf import render_report_pdf

pytestmark = pytest.mark.django_db

TODAY = date(2026, 9, 9)


# ── La période précédente ──────────────────────────────────────────────────


def test_a_named_month_compares_to_the_same_month_last_year():
    """Septembre 2026 se compare à septembre 2025, pas à août 2026.

    Une bibliothèque ne fait pas le même mois en pleine rentrée et en plein
    été ; d'une année sur l'autre, si.
    """
    september = periods.named_period("month", TODAY)
    before = periods.previous_equivalent(september)
    assert (before.start, before.end) == (date(2025, 9, 1), date(2025, 9, 30))
    assert before.label == september.label.replace("2026", "2025")


def test_the_previous_month_also_steps_back_a_year():
    august = periods.named_period("prev_month", TODAY)
    before = periods.previous_equivalent(august)
    assert (before.start, before.end) == (date(2025, 8, 1), date(2025, 8, 31))


def test_a_named_year_compares_to_the_year_before():
    for key, expected in (("year", 2025), ("prev_year", 2024)):
        before = periods.previous_equivalent(periods.named_period(key, TODAY))
        assert (before.start, before.end) == (date(expected, 1, 1), date(expected, 12, 31))
        assert before.label == str(expected)


def test_the_last_twelve_months_compare_to_the_twelve_before():
    """Val : « 12 derniers mois » → du 10/09/2024 au 09/09/2025."""
    current = periods.default_period(TODAY)
    before = periods.previous_equivalent(current)
    assert (before.start, before.end) == (date(2024, 9, 10), date(2025, 9, 9))
    assert before.label == "10/09/24 – 09/09/25"


def test_free_dates_compare_to_the_block_that_precedes_them():
    """Le bloc de même durée qui précède, sans chevauchement.

    On ne peut pas décaler d'un an une fenêtre dont on ignore l'intention. Et
    le bloc s'arrête la **veille** du début : la formule de Val, prise au pied
    de la lettre, faisait compter le premier jour des deux côtés.
    """
    current = periods.Period(date(2026, 3, 1), date(2026, 3, 11), "libre", "custom")
    before = periods.previous_equivalent(current)
    assert (before.start, before.end) == (date(2026, 2, 18), date(2026, 2, 28))
    assert before.days == current.days
    assert before.end < current.start
    assert before.label == "18/02/26 – 28/02/26"


def test_a_leap_day_steps_back_without_crashing():
    """`date.replace(year=…)` lève sur un 29 février ; le mois comparé reste
    février, la retombée au 28 ne change aucun résultat."""
    february = periods.named_period("month", date(2024, 2, 29))
    before = periods.previous_equivalent(february)
    assert (before.start, before.end) == (date(2023, 2, 1), date(2023, 2, 28))


# ── Le bilan ───────────────────────────────────────────────────────────────


@pytest.fixture
def dataset(django_user_model):
    category = Category.objects.create(code="ALB", name="Albums")
    other = Category.objects.create(code="DOC", name="Documentaires")
    mcat = MemberCategory.objects.create(code="AD", name="Adulte")
    member = Member.objects.create(first_name="Ada", last_name="Lovelace", category=mcat)

    record = BibliographicRecord.objects.create(title="Fondation", category=category)
    doc = BibliographicRecord.objects.create(title="Atlas", category=other)
    item = Item.objects.create(record=record)
    Item.objects.create(record=record)
    Item.objects.create(record=doc)

    Loan.objects.create(
        item=item, member=member, due_date=date.today() + timedelta(days=21),
        status=LoanStatus.ACTIVE,
    )
    return {"member": member, "item": item, "category": category}


def _period():
    return periods.Period(
        date.today() - timedelta(days=60), date.today(), "test", "custom"
    )


def test_the_summary_table_has_a_previous_period_column(dataset):
    period = periods.named_period("year")
    page = builders.get("overview").builder(period, {})
    table = next(t for t in page.tables if t.title == "Le bilan en un tableau")

    assert len(table.columns) == 3
    assert table.columns[1].label == period.label
    assert table.columns[2].label == periods.previous_equivalent(period).label
    assert all(len(row) == 3 for row in table.rows)


def test_todays_figures_leave_the_comparison_column_empty(dataset):
    """Les huit premières lignes décrivent aujourd'hui : rien à comparer."""
    page = builders.get("overview").builder(periods.named_period("year"), {})
    table = next(t for t in page.tables if t.title == "Le bilan en un tableau")
    for row in table.rows[:8]:
        assert row[2] == ""
    for row in table.rows[8:]:
        assert row[2] != ""


def test_the_summary_table_has_no_chart_any_more(dataset):
    """Quatorze grandeurs hétérogènes ne se lisent pas sur un axe commun."""
    page = builders.get("overview").builder(_period(), {})
    table = next(t for t in page.tables if t.title == "Le bilan en un tableau")
    assert table.chart is None


# ── Le double histogramme des rayons ───────────────────────────────────────


def test_the_sections_table_carries_a_two_series_bar_chart(dataset):
    page = builders.get("collection").builder(_period(), {})
    table = next(t for t in page.tables if t.title.startswith("Les rayons"))
    chart = table.chart

    assert chart is not None
    assert chart.kind == "bar"
    assert [s.label for s in chart.series] == ["Livres", "Prêts sur la période"]
    # Une valeur par rayon dans chaque série, alignée sur les libellés.
    assert all(len(s.values) == len(chart.labels) for s in chart.series)
    assert not chart.is_empty


def test_the_bar_chart_figures_match_the_table(dataset):
    """Le graphe montre les mêmes chiffres que les colonnes du tableau."""
    page = builders.get("collection").builder(_period(), {})
    table = next(t for t in page.tables if t.title.startswith("Les rayons"))
    books, loans = table.chart.series
    for index, label in enumerate(table.chart.labels):
        row = next(r for r in table.rows if r[0] == label)
        assert str(books.values[index]) == row[1]
        assert str(loans.values[index]) == row[3]


def test_a_grouped_bar_svg_draws_one_bar_per_series_and_a_legend():
    chart = Chart(
        title="t", kind="bar", labels=["Albums", "Romans"],
        series=[Series("Livres", [40, 12]), Series("Prêts", [7, 3])],
    )
    svg = charts.bar_svg(chart)
    # Deux rayons × deux séries, plus deux pastilles de légende.
    assert svg.count("<rect") == 6
    assert "Livres" in svg and "Prêts" in svg


def test_a_single_series_bar_chart_still_works():
    """Un histogramme simple est le cas particulier d'un groupé à une série :
    il ne doit pas avoir de légende, ni de barre en trop."""
    chart = Chart(title="t", kind="bar", labels=["A", "B", "C"], values=[3, 1, 2])
    svg = charts.bar_svg(chart)
    assert svg.count("<rect") == 3


def test_a_grouped_bar_chart_with_only_zeros_is_empty():
    chart = Chart(
        title="t", kind="bar", labels=["A"],
        series=[Series("Livres", [0]), Series("Prêts", [0])],
    )
    assert chart.is_empty


def test_the_grouped_chart_survives_pdf_and_excel(dataset):
    """Les trois moteurs doivent savoir grouper, pas seulement le SVG."""
    page = builders.get("collection").builder(_period(), {})
    table = next(t for t in page.tables if t.title.startswith("Les rayons"))

    pdf = render_report_pdf(page, table)
    assert pdf[:4] == b"%PDF"

    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(render_report_xlsx(page, table)))
    graph = workbook[workbook.sheetnames[-1]]._charts[0]
    assert len(graph.series) == 2
    assert graph.grouping == "clustered"


# ── Le guide utilisateur ───────────────────────────────────────────────────


def test_the_guide_page_lists_every_report_and_sub_report(dataset):
    """« Tous les rapports » doit rester en phase avec les écrans réels.

    La page est produite par `scripts/build_reports_guide.py` : chaque ancre est
    l'empreinte du titre d'un sous-rapport. Renommer un tableau sans rejouer le
    script casserait le lien **en silence** — le guide continuerait de pointer
    sur une ancre qui n'existe plus.
    """
    import re
    from pathlib import Path

    from django.conf import settings
    from django.utils import translation

    page = Path(settings.BASE_DIR) / "docs/user-guide/docs/rapports/liste.md"
    assert page.exists(), "lancer `python scripts/build_reports_guide.py`"
    text = page.read_text(encoding="utf-8")

    with translation.override("fr"):
        period = periods.default_period()
        for spec in builders.REPORTS:
            url = f"/bibliofelia/fr/reports/{spec.slug}/"
            assert url in text, f"l'écran « {spec.slug} » manque au guide"

            built = spec.builder(period if spec.needs_period else None, {})
            for block in built.blocks:
                anchor = f"{url}#{block.key}"
                assert anchor in text, (
                    f"« {block.title} » ({spec.slug}) manque au guide — "
                    "rejouer scripts/build_reports_guide.py"
                )

        # Et l'inverse : aucune ancre du guide ne pointe dans le vide.
        known = set()
        for spec in builders.REPORTS:
            built = spec.builder(period if spec.needs_period else None, {})
            known.update(
                f"/bibliofelia/fr/reports/{spec.slug}/#{b.key}" for b in built.blocks
            )
        for link in re.findall(r"/bibliofelia/fr/reports/[a-z-]+/#\w+", text):
            assert link in known, f"{link} ne correspond à aucun sous-rapport"


def test_the_guide_exists_in_the_four_languages():
    from pathlib import Path

    from django.conf import settings

    guide = Path(settings.BASE_DIR) / "docs/user-guide/docs/rapports"
    for suffix in ("", ".en", ".es", ".mg"):
        assert (guide / f"liste{suffix}.md").exists()
